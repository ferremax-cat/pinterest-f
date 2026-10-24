#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador del índice de búsqueda v2 — Catálogo Ferremax

Lee json/productos.json y escribe json/search/v2/indice.json.
Antes de generar verifica los casos de js/search/v2/casos_normalizacion.json;
si alguno falla, no escribe nada (así un error nunca publica un índice roto).

Uso, desde cualquier carpeta del proyecto:
    python generar_indice_busqueda.py                      genera el índice
    python generar_indice_busqueda.py --probar             solo verifica los casos
    python generar_indice_busqueda.py --consulta "fix 3.5" muestra tokens y resultados

La normalización está escrita sin "lookbehind" en las expresiones regulares
para poder copiarla tal cual al motor en JavaScript. Las dos versiones deben
dar exactamente los mismos tokens: eso es lo que controlan los casos de prueba.
"""
import argparse
import gzip
import hashlib
import json
import re
import sys
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

# Antes de pisar el indice publicado, comparar la cantidad de productos
from scripts.guardia_registros import revisar, contar_json, EscrituraBloqueada


# ---------------------------------------------------------------------------
# Rutas: se busca la raíz del proyecto subiendo hasta encontrar json/productos.json
# ---------------------------------------------------------------------------
def encontrar_raiz():
    carpeta = Path(__file__).resolve().parent
    for candidata in [carpeta, *carpeta.parents]:
        if (candidata / "json" / "productos.json").exists():
            return candidata
    sys.exit("No se encontró json/productos.json subiendo desde " + str(carpeta))


RAIZ = encontrar_raiz()
RUTA_PRODUCTOS = RAIZ / "json" / "productos.json"
RUTA_SALIDA = RAIZ / "json" / "search" / "v2" / "indice.json"
RUTA_CASOS = RAIZ / "js" / "search" / "v2" / "casos_normalizacion.json"
RUTA_VERSION = RAIZ / "json" / "version.json"
# Diccionario de equivalencias, editable con Excel (una fila por grupo)
RUTA_DICCIONARIO = RAIZ / "excel" / "diccionario_buscador.xlsx"
# Para revisar en Excel; salida/ está en .gitignore: no se commitea ni se publica
RUTA_CANDIDATOS = RAIZ / "salida" / "candidatos_diccionario.xlsx"


# ---------------------------------------------------------------------------
# Normalización y tokenización (debe ser idéntica en JavaScript)
# ---------------------------------------------------------------------------
PALABRAS_VACIAS = {"de", "del", "la", "las", "el", "los", "para", "con",
                   "en", "por", "al", "un", "una"}

RE_ACENTOS = re.compile(r"[\u0300-\u036f]")                  # tildes y la de la ñ
# Siglas con puntos: r.c.t / r.c.t. -> rct. Solo letras sueltas: no toca decimales
# (3.5), fracciones ni abreviaturas de más de una letra (valv.esf)
RE_SIGLA = re.compile(r"(^|[^a-z0-9.])((?:[a-z]\.)+[a-z])\.?(?![a-z0-9])")
# Letras sueltas con guion: h-h, m-h, h-h-h -> h/h, m/h, h/h/h (como M/M). No toca
# FUS-FUS (más de una letra), 20-32 ni T-6 (números)
RE_LETRAS_GUION = re.compile(r"(^|[^a-z0-9-])([a-z](?:-[a-z])+)(?![a-z0-9-])")
RE_COMA_DECIMAL = re.compile(r"(\d),(?=\d)")                 # 3,5  -> 3.5
RE_CEROS_FINALES = re.compile(r"(\d+\.\d*?[1-9])0+(?!\d)")   # 3.50 -> 3.5
RE_DECIMAL_CERO = re.compile(r"(\d+)\.0+(?!\d)")             # 4.0  -> 4
RE_X_ENTRE_MEDIDAS = re.compile(r"(\d[a-z\"]{0,3})x(?=\d)")  # 20mmx4mt -> 20mm x 4mt
RE_X_INICIAL = re.compile(r"(^|\s)x(?=\d)")                  # x15m -> x 15m
RE_FRACCION_MIXTA = re.compile(r"(^|\s)(\d{1,2})[\s.\-]+(\d+/\d+)")  # 1 1/2, 1.1/2 -> 1-1/2
RE_FRACCION_PEGADA = re.compile(r"(^|[^\d/.])(\d)(\d)/(\d{1,2})(?![\d/])")  # 11/2 -> 1-1/2
DENOMINADORES_PULGADA = {2, 4, 8, 16}


def _separar_fraccion_pegada(m):
    entero, numerador, denominador = int(m.group(2)), int(m.group(3)), int(m.group(4))
    if (denominador in DENOMINADORES_PULGADA and numerador < denominador
            and entero * 10 + numerador > denominador):
        return m.group(1) + m.group(2) + "-" + m.group(3) + "/" + m.group(4)
    return m.group(0)
RE_TOKEN = re.compile(
    r"\d+-\d+/\d+"                # 1-1/2 (fracción mixta)
    r"|\d+(?:[./]\d+)+[a-z]*"     # 3.5  1/2  3.5mm
    r"|\d+[a-z]+\d*[a-z]*"        # 20mm  15m  4mt
    r"|\d+"                       # 20
    r"|[a-z](?:/[a-z])+(?![a-z])" # f/f  i/n  h/h/h (no corta C/TUERCA en c/t + uerca)
    r"|[a-z]+\d+[a-z\d]*"         # pn16
    r"|[a-z]+"                    # palabras
)
RE_PARTES = re.compile(r"\d+(?:[./]\d+)*|[a-z]+")
RE_TIENE_DIGITO = re.compile(r"\d")
RE_TIENE_LETRA = re.compile(r"[a-z]")
RE_PREFIJO_CATEGORIA = re.compile(r"^\s*\d+\s*\.\s*")       # "8.GUANTES" -> "GUANTES"


def normalizar(texto):
    t = str(texto or "").lower()
    t = unicodedata.normalize("NFD", t)
    t = RE_ACENTOS.sub("", t)
    t = RE_SIGLA.sub(lambda m: m.group(1) + m.group(2).replace(".", ""), t)
    t = RE_LETRAS_GUION.sub(lambda m: m.group(1) + m.group(2).replace("-", "/"), t)
    t = RE_COMA_DECIMAL.sub(r"\1.", t)
    t = RE_CEROS_FINALES.sub(r"\1", t)
    t = RE_DECIMAL_CERO.sub(r"\1", t)
    t = RE_FRACCION_MIXTA.sub(r"\1\2-\3", t)
    t = RE_FRACCION_PEGADA.sub(_separar_fraccion_pegada, t)
    t = RE_X_ENTRE_MEDIDAS.sub(r"\1 x ", t)
    t = RE_X_INICIAL.sub(r"\1x ", t)
    return t


def tokenizar(texto, expandir=True):
    """expandir=True para el índice (20mm también se indexa como 20 y mm).
    expandir=False para la consulta del usuario."""
    salida, vistos = [], set()

    def agregar(tok):
        if tok in vistos:
            return
        if tok.isalpha() and (len(tok) < 2 or tok in PALABRAS_VACIAS):
            return
        vistos.add(tok)
        salida.append(tok)

    for m in RE_TOKEN.finditer(normalizar(texto)):
        tok = m.group(0)
        agregar(tok)
        if expandir and RE_TIENE_DIGITO.search(tok) and RE_TIENE_LETRA.search(tok):
            # 20mm -> 20 y mm. Pero pn20 -> solo pn: el 20 de una clase o modelo
            # no es una medida y haría que "20" encuentre caños de 25, 32, etc.
            empieza_con_letra = tok[0].isalpha()
            for parte in RE_PARTES.findall(tok):
                if empieza_con_letra and parte[0].isdigit():
                    continue
                agregar(parte)
    return salida


# ---------------------------------------------------------------------------
# Orden natural: los ids se asignan en este orden, y sirven de desempate
# (3.5x16, 3.5x20, 3.5x25 en lugar del orden del archivo)
# ---------------------------------------------------------------------------
def valor_numerico(txt):
    try:
        if "/" in txt:
            a, b = txt.split("/", 1)
            return float(a) / float(b) if float(b) else 0.0
        return float(txt)
    except ValueError:
        return 0.0


def clave_natural(nombre_normalizado):
    clave = []
    for parte in re.findall(r"\d+(?:[./]\d+)?|[a-z]+", nombre_normalizado):
        if parte[0].isdigit():
            clave.append((0, valor_numerico(parte), ""))
        else:
            clave.append((1, 0.0, parte))
    return clave


# ---------------------------------------------------------------------------
# Construcción del índice
# ---------------------------------------------------------------------------
def cargar_openpyxl():
    """openpyxl se usa solo para el diccionario y los candidatos (Excel). Si no
    está instalado, se explica cómo instalarlo en lugar de un error de Python."""
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        sys.exit("Falta la librería openpyxl, que se usa para leer el diccionario (Excel).\n"
                 "Instalala con:  python -m pip install openpyxl")


def dar_formato(hoja, anchos):
    """Planilla cómoda: encabezados en negrita, primera fila fija, filtro y
    columnas con el ancho indicado."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    for i, ancho in enumerate(anchos, 1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho


def leer_diccionario():
    """Grupos de términos equivalentes de excel/diccionario_buscador.xlsx, ya
    normalizados. Primera hoja; la fila 1 tiene los encabezados: "principal",
    "variante 1", "variante 2"... y "comentario", que se ignora (igual que
    cualquier otra columna). Una fila por grupo; las filas vacías se saltean.
    Devuelve (grupos, errores), con el número de fila de Excel en cada error.
    Un error frena la generación, igual que un caso de normalización fallido."""
    if not RUTA_DICCIONARIO.exists():
        print("AVISO: no existe " + str(RUTA_DICCIONARIO.relative_to(RAIZ)) + "; índice sin equivalencias")
        return [], []
    openpyxl = cargar_openpyxl()
    try:
        libro = openpyxl.load_workbook(RUTA_DICCIONARIO, read_only=True, data_only=True)
    except Exception as e:
        return [], ["no se pudo abrir {}: {}".format(RUTA_DICCIONARIO.name, e)]
    try:
        filas = libro.worksheets[0].iter_rows(values_only=True)
        encabezados = [str(c or "").strip().lower() for c in next(filas, ())]
        if not encabezados or encabezados[0] != "principal":
            return [], ["fila 1: la primera columna tiene que llamarse 'principal'"]
        # principal y las variantes, en el orden de las columnas
        columnas = [i for i, e in enumerate(encabezados) if i == 0 or e.startswith("variante")]

        grupos, errores, fila_de = [], [], {}
        for nro, fila in enumerate(filas, 2):
            celdas = [str(fila[i]).strip() if i < len(fila) and fila[i] is not None else ""
                      for i in columnas]
            if not any(celdas):
                continue                   # fila vacía
            if not celdas[0]:
                errores.append("fila {}: falta la palabra principal".format(nro))
                continue
            terminos = []
            for celda in celdas:
                if not celda:
                    continue
                termino = normalizar(celda).strip()
                if tokenizar(celda, expandir=False) != [termino]:
                    errores.append("fila {}: '{}' no es una sola palabra válida".format(nro, celda))
                    continue
                if termino not in terminos:
                    terminos.append(termino)
            if not terminos:
                continue
            if len(terminos) == 1:
                errores.append("fila {}: el grupo '{}' tiene un solo término".format(nro, terminos[0]))
            for t in terminos:
                if t in fila_de:
                    errores.append("fila {}: '{}' ya está en el grupo de la fila {}".format(nro, t, fila_de[t]))
                else:
                    fila_de[t] = nro
            grupos.append(terminos)
        return grupos, errores
    finally:
        libro.close()


def cargar_productos():
    with open(RUTA_PRODUCTOS, encoding="utf-8") as f:
        datos = json.load(f)
    if not isinstance(datos, dict):
        sys.exit("Formato inesperado en productos.json: se esperaba un objeto indexado por código")
    return {cod: p for cod, p in datos.items() if isinstance(p, dict)}


def construir_indice(grupos=()):
    filas = []
    for codigo, p in cargar_productos().items():
        nombre = p.get("name") or p.get("nombre") or ""
        categoria = p.get("category") or p.get("categoria") or ""
        filas.append((clave_natural(normalizar(nombre)), str(codigo), nombre, categoria))
    filas.sort(key=lambda f: (f[0], f[1]))

    # Cada término del diccionario indexa también a todos los de su grupo:
    # un producto que dice "BCE" se encuentra por "bronce" y viceversa
    grupo_de = {t: g for g in grupos for t in g}

    def con_equivalentes(tokens):
        salida, vistos = list(tokens), set(tokens)
        for tok in tokens:
            for eq in grupo_de.get(tok, ()):
                if eq not in vistos:
                    vistos.add(eq)
                    salida.append(eq)
        return salida

    productos, ix_nombre, ix_categoria = [], {}, {}
    for i, (_, codigo, nombre, categoria) in enumerate(filas):
        productos.append([codigo, nombre, categoria])
        for tok in con_equivalentes(tokenizar(nombre)):
            ix_nombre.setdefault(tok, []).append(i)
        for tok in con_equivalentes(tokenizar(RE_PREFIJO_CATEGORIA.sub("", categoria))):
            ix_categoria.setdefault(tok, []).append(i)

    # Huella del contenido: cambia solo si cambian códigos, nombres o categorías
    # (no los precios). Con eso el workflow, que corre cada hora, no commitea un
    # índice igual al anterior
    cuerpo = json.dumps([productos, ix_nombre, ix_categoria, list(grupos)], ensure_ascii=False, separators=(",", ":"))
    huella = hashlib.sha256(cuerpo.encode("utf-8")).hexdigest()[:16]
    return {
        "version": "v2-" + huella,
        "generado": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "total": len(productos),
        "productos": productos,
        "nombre": ix_nombre,
        "categoria": ix_categoria,
        # Para el navegador: qué términos son equivalentes (el primero de cada
        # grupo es el que se usa para ordenar). Unos 2 KB
        "equivalencias": [list(g) for g in grupos],
    }


# ---------------------------------------------------------------------------
# Casos de prueba y consulta de verificación
# ---------------------------------------------------------------------------
def probar_casos():
    if not RUTA_CASOS.exists():
        print("AVISO: no existe " + str(RUTA_CASOS.relative_to(RAIZ)))
        return False
    with open(RUTA_CASOS, encoding="utf-8") as f:
        casos = json.load(f)["casos"]
    fallas = 0
    for c in casos:
        obtenido = tokenizar(c["texto"], c.get("expandir", True))
        if obtenido != c["tokens"]:
            fallas += 1
            print("FALLA  " + repr(c["texto"]))
            print("   esperado: " + str(c["tokens"]))
            print("   obtenido: " + str(obtenido))
    print("Casos de normalización: {} de {} correctos".format(len(casos) - fallas, len(casos)))
    return fallas == 0


def consultar(indice, consulta, limite=15):
    tokens = tokenizar(consulta, expandir=False)
    print("Tokens de la consulta: " + str(tokens))
    if not tokens:
        return
    ix = indice["nombre"]
    conjuntos = [set(ix.get(t, [])) for t in tokens[:-1]]
    ultimo = tokens[-1]
    por_prefijo = set()
    for tok, ids in ix.items():
        if tok.startswith(ultimo):
            por_prefijo.update(ids)
    conjuntos.append(por_prefijo)
    resultado = sorted(set.intersection(*conjuntos))
    print("Productos con todas las palabras: {}".format(len(resultado)))
    for i in resultado[:limite]:
        codigo, nombre, _ = indice["productos"][i]
        print("   {:<16} {}".format(codigo, nombre))


def escribir_candidatos(grupos):
    """Palabras del catálogo que no están en el diccionario y parecen abreviaturas,
    para revisar en Excel: de hasta 5 letras con una vocal como mucho (bce, ngro,
    std), o el comienzo de otra palabra más larga del catálogo (galv, desmalez).
    Con la cantidad de productos, las palabras largas que empiezan igual y tres
    nombres de ejemplo, de más a menos productos."""
    en_dic = {t for g in grupos for t in g}
    nombres_de = {}
    for p in cargar_productos().values():
        nombre = (p.get("name") or p.get("nombre") or "").strip()
        for tok in set(tokenizar(nombre, expandir=False)):
            if tok.isalpha():
                nombres_de.setdefault(tok, []).append(nombre)
    palabras = sorted(nombres_de)
    filas = []
    for i, tok in enumerate(palabras):
        if tok in en_dic or len(nombres_de[tok]) < 2:
            continue
        largas = []
        for otra in palabras[i + 1:]:
            if not otra.startswith(tok):
                break
            if len(otra) >= len(tok) + 2:
                largas.append(otra)
        vocales = sum(1 for ch in tok if ch in "aeiou")
        if not ((len(tok) <= 5 and vocales <= 1) or (len(tok) >= 3 and largas)):
            continue
        largas.sort(key=lambda w: -len(nombres_de[w]))
        ejemplos = (nombres_de[tok] + ["", "", ""])[:3]
        filas.append([tok, len(nombres_de[tok]), ", ".join(largas[:3])] + ejemplos)
    filas.sort(key=lambda f: -f[1])

    openpyxl = cargar_openpyxl()
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "candidatos"
    hoja.append(["palabra", "productos", "palabras que empiezan igual", "ejemplo 1", "ejemplo 2", "ejemplo 3"])
    for fila in filas:
        hoja.append(fila)
    dar_formato(hoja, [14, 11, 36, 46, 46, 46])
    RUTA_CANDIDATOS.parent.mkdir(parents=True, exist_ok=True)
    libro.save(RUTA_CANDIDATOS)
    print("Candidatos: {} palabras en {}".format(len(filas), RUTA_CANDIDATOS.relative_to(RAIZ)))


def grupos_de_indice(ruta):
    """Grupos de equivalencias guardados en un índice ya generado (el que estaba
    publicado): sirven para rearmar el índice de hoy con el diccionario anterior."""
    try:
        with open(ruta, encoding="utf-8") as f:
            return [list(g) for g in json.load(f).get("equivalencias", [])]
    except (OSError, ValueError) as e:
        sys.exit("No se pudo leer el diccionario de {}: {}".format(ruta, e))


def leer_version_anterior():
    """Versión del índice ya publicado, o None si no existe o no se puede leer."""
    try:
        with open(RUTA_SALIDA, encoding="utf-8") as f:
            return json.load(f).get("version")
    except (OSError, ValueError):
        return None


def registrar_version(version):
    """Deja en json/version.json la huella del índice publicado, para que las
    pestañas abiertas lo actualicen sin recargar. Conserva las demás claves
    (productos, que escribe excel_to_json.py) y solo escribe si cambió."""
    datos = {}
    if RUTA_VERSION.exists():
        try:
            with open(RUTA_VERSION, encoding="utf-8") as f:
                datos = json.load(f)
        except (OSError, ValueError) as e:
            # No se reescribe: se perdería la huella de productos
            print("AVISO: no se pudo leer {} ({}); no se registra la versión del índice".format(
                RUTA_VERSION.relative_to(RAIZ), e))
            return
    if datos.get("indice", {}).get("huella") == version:
        return
    datos["indice"] = {"huella": version}
    # Mismo formato que excel_to_json.py (indent=2, sin salto final)
    temporal = RUTA_VERSION.with_suffix(".tmp")
    temporal.write_text(json.dumps(datos, indent=2), encoding="utf-8")
    temporal.replace(RUTA_VERSION)
    print("version.json actualizado: indice " + version)


# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="Índice de búsqueda v2 — Ferremax")
    ap.add_argument("--probar", action="store_true", help="solo verificar los casos")
    ap.add_argument("--consulta", help="mostrar tokens y resultados de una búsqueda")
    ap.add_argument("--candidatos", action="store_true",
                    help="escribir salida/candidatos_diccionario.xlsx para revisar en Excel")
    ap.add_argument("--grupos-de", metavar="INDICE",
                    help="usar las equivalencias de un índice ya generado en lugar del Excel "
                         "(el workflow lo usa para volver al diccionario anterior)")
    ap.add_argument("--salida", metavar="RUTA",
                    help="escribir el índice en otra ruta, sin tocar el publicado ni version.json")
    args = ap.parse_args()

    if not probar_casos():
        sys.exit("Hay casos de normalización fallidos: no se genera el índice.")
    if args.grupos_de:
        grupos = grupos_de_indice(args.grupos_de)
        print("Diccionario tomado de {}: {} grupos".format(args.grupos_de, len(grupos)))
    else:
        grupos, errores = leer_diccionario()
        for e in errores:
            print("ERROR diccionario_buscador.xlsx, " + e)
        if errores:
            sys.exit("El diccionario tiene errores: no se genera el índice.")
        print("Diccionario: {} grupos, {} términos".format(len(grupos), sum(len(g) for g in grupos)))
    if args.probar:
        return
    if args.candidatos:
        escribir_candidatos(grupos)
        return

    indice = construir_indice(grupos)

    if args.consulta:
        consultar(indice, args.consulta)
        return

    # Índice de prueba: no toca el publicado ni version.json
    if args.salida:
        destino = Path(args.salida)
        destino.parent.mkdir(parents=True, exist_ok=True)
        destino.write_text(json.dumps(indice, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        print("Índice de prueba: " + str(destino))
        return

    # Mismo contenido que el publicado: no se toca el índice (no hay commit), pero
    # version.json igual tiene que tener su huella
    if leer_version_anterior() == indice["version"]:
        print("Índice sin cambios (" + indice["version"] + "): no se reescribe")
        registrar_version(indice["version"])
        return

    # Antes de pisar el publicado: con muchos menos productos que el anterior,
    # productos.json vino incompleto (ver scripts/guardia_registros.py)
    try:
        revisar(RUTA_SALIDA.name, contar_json(RUTA_SALIDA, lambda d: len(d["productos"])), indice["total"])
    except EscrituraBloqueada as e:
        sys.exit(str(e))

    contenido = json.dumps(indice, ensure_ascii=False, separators=(",", ":"))
    RUTA_SALIDA.parent.mkdir(parents=True, exist_ok=True)
    temporal = RUTA_SALIDA.with_suffix(".tmp")
    temporal.write_text(contenido, encoding="utf-8")
    temporal.replace(RUTA_SALIDA)
    # Después del índice: la huella nunca anuncia un archivo que todavía no existe
    registrar_version(indice["version"])

    bytes_total = len(contenido.encode("utf-8"))
    bytes_gzip = len(gzip.compress(contenido.encode("utf-8")))
    print("Índice generado: " + str(RUTA_SALIDA.relative_to(RAIZ)))
    print("   versión:          " + indice["version"])
    print("   productos:        {}".format(indice["total"]))
    print("   tokens de nombre: {}".format(len(indice["nombre"])))
    print("   tokens categoría: {}".format(len(indice["categoria"])))
    print("   tamaño:           {:.2f} MB ({:.2f} MB comprimido)".format(
        bytes_total / 1048576, bytes_gzip / 1048576))


if __name__ == "__main__":
    main()
